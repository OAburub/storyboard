"""
PPTX Engine — LectureBuilder (Thin Composer)
=============================================

Composes modular capabilities into a single builder class.

Architecture:
    _pptx_core.py          — SlideEngine base + constants + helpers
    _pptx_depth.py         — Visual depth system (washes, accents, corners)
    _pptx_structural.py    — Structural slides (title, objectives, divider, summary, closing)
    _pptx_interactions.py  — Interactive slides (quiz, drag-drop, click-reveal, scenario)
    pptx_engine.py         — THIS FILE: LectureBuilder composer (public API)

The LectureBuilder class inherits from all modules via Python's MRO
(Method Resolution Order). Mixins come BEFORE SlideEngine so their
methods can override or extend base behavior if needed.

Usage:
    from pptx_engine import LectureBuilder

    builder = LectureBuilder(
        project_code="DSAI",
        unit_number=1,
        unit_name="المهارات الرقمية",
        institution="جامعة نجران - كلية علوم الحاسب ونظم المعلومات"
    )
    builder.add_title_slide(title="المحاضرة الأولى", subtitle="مقدمة في علوم الحاسوب")
    builder.add_objectives_slide(objectives=["تعريف ماهي التقنية", "التعرف على الفوائد"])
    builder.add_content_slide(title="المقدمة", bullets=["نقطة أولى", "نقطة ثانية"])
    builder.save("output/DSAI/U01/DSAI_U01_Interactive_Lecture.pptx")
"""

import os
from typing import Optional

from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

# --- Module imports ---
# SlideEngine: base class with __init__, save(), finalize(), all private helpers
# All design constants (colors, fonts, positions, assets) are re-exported here
# so that `from pptx_engine import LectureBuilder, PRIMARY_BLUE` still works.
from _pptx_core import (
    SlideEngine,
    # Slide dimensions
    SLIDE_WIDTH,
    SLIDE_HEIGHT,
    # Color palette
    PRIMARY_BLUE,
    ACCENT1_BLUE,
    BODY_TEXT,
    SUBTITLE_TEXT,
    LINK_BLUE,
    WHITE,
    DARK_BG,
    BUTTON_BORDER,
    NOTES_YELLOW,
    TEAL,
    ACCENT_GREEN,
    ACCENT_RED,
    ACCENT_ORANGE,
    LIGHT_BLUE_BG,
    CONTENT_CARD_BG,
    CONTENT_CARD_BORDER,
    CARD_LIGHT_BG,
    OPTION_ALT_BG,
    DIVIDER_BG,
    BULLET_MARKER_COLOR,
    SHADOW_COLOR,
    PRIMARY_BLUE_LIGHT,
    PRIMARY_BLUE_DARK,
    TEAL_LIGHT,
    WARM_GRAY,
    ACCENT_DEFINITION,
    ACCENT_EXAMPLE,
    HEADER_BAR_BLUE,
    ACCENT_CYCLE,
    # PNG asset names
    ASSET_BANNER_NARROW,
    ASSET_BANNER_WIDE,
    ASSET_OBJECTIVE_ROW,
    ASSET_TARGET_ICON,
    ASSET_PLAY_ICON,
    ASSET_HAND_CURSOR,
    ASSET_CORNER_TR,
    ASSET_CORNER_BL,
    ASSET_TEXT_BUBBLE,
    ASSET_IMAGE_FRAME,
    ASSET_CONTENT_BG,
    # Font names
    FONT_EXTRABOLD,
    FONT_MEDIUM,
    FONT_REGULAR,
    FONT_FALLBACK,
    # Positions
    TITLE_BAR_LEFT,
    TITLE_BAR_TOP,
    TITLE_BAR_WIDTH,
    TITLE_BAR_HEIGHT,
    BANNER_LEFT,
    BANNER_TOP,
    BANNER_WIDTH,
    BANNER_HEIGHT,
    WIDE_BANNER_LEFT,
    WIDE_BANNER_TOP,
    WIDE_BANNER_WIDTH,
    WIDE_BANNER_HEIGHT,
    WIDE_BANNER_TEXT_LEFT,
    WIDE_BANNER_TEXT_TOP,
    WIDE_BANNER_TEXT_WIDTH,
    WIDE_BANNER_TEXT_HEIGHT,
    NARROW_BANNER_TEXT_LEFT,
    NARROW_BANNER_TEXT_TOP,
    NARROW_BANNER_TEXT_WIDTH,
    NARROW_BANNER_TEXT_HEIGHT,
    CONTENT_LEFT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    CONTENT_HEIGHT,
    TEXT_MARGIN_LR,
    TEXT_MARGIN_TB,
    # Paths
    TEMPLATE_PATH,
)

# DepthMixin: visual depth (washes, accents, corners, progress dots)
from _pptx_depth import DepthMixin

# StructuralMixin: structural slides (title, objectives, section divider, summary, closing)
from _pptx_structural import StructuralMixin

# VisualGrammarMixin: rich content slides (process flow, stat cards, quote, timeline, etc.)
from _pptx_visual_grammar import VisualGrammarMixin

# InteractionsMixin: interactive slides (quiz, drag-drop, click-reveal, slider, dropdown, scenario)
from _pptx_interactions import InteractionsMixin


class LectureBuilder(StructuralMixin, InteractionsMixin, VisualGrammarMixin, DepthMixin, SlideEngine):
    """
    Production PPTX builder for Arabic e-learning lectures.

    Composes:
        SlideEngine         (core)         — presentation setup, helpers, save/finalize
        DepthMixin          (depth)        — visual depth washes, accents, decorative corners
        VisualGrammarMixin  (content)      — rich visual patterns (process flow, stat cards, etc.)
        InteractionsMixin   (interactions) — quiz, drag-drop, click-reveal, slider, dropdown, scenario
        StructuralMixin     (frame)        — title, objectives, section divider, summary, closing

    Usage:
        builder = LectureBuilder(
            project_code="CODE", unit_number=1,
            unit_name="اسم الوحدة", institution="اسم الجامعة"
        )
        builder.add_title_slide(title="المحاضرة الأولى:", subtitle="مقدمة")
        builder.add_content_slide(title="المحتوى", bullets=["نقطة 1", "نقطة 2"])
        builder.save("output/CODE/U01/lecture.pptx")
    """

    def __init__(self, project_code, unit_number, unit_name, institution,
                 designer="", template_path=None):
        super().__init__(project_code, unit_number, unit_name, institution,
                         designer, template_path)

    # ==========================================================================
    # TEMPORARY: Content slide methods (will move to _pptx_visual_grammar.py)
    # ==========================================================================

    def add_content_slide(
        self,
        title: str,
        bullets: Optional[list] = None,
        paragraphs: Optional[list] = None,
        image_placeholder: Optional[str] = None,
        image_path: Optional[str] = None,
        image_prompt: Optional[str] = None,
        notes: str = "",
    ):
        """
        Add a content slide with a header and body text.

        This is the main workhorse slide type — presents topic content
        with a title banner and body text. Optionally includes an image.

        Args:
            title: Section title (e.g., "المقدمة")
            bullets: List of bullet point strings (use this OR paragraphs)
            paragraphs: List of paragraph strings (use this OR bullets)
            image_placeholder: Optional text describing what image to add
            image_path: Optional path to a real image file (PNG/JPG)
            image_prompt: Optional prompt to auto-generate an image
            notes: Speaker notes / Storyline instructions
        """
        self.slide_count += 1
        slide = self._add_content_slide_with_layout()

        # Set hidden TOC title for Storyline sidebar menu
        self._set_slide_title_for_toc(slide, title)

        # --- Lecture title bar at top ---
        self._add_header_bar(slide, self.lecture_title)

        # --- Section banner ---
        self._add_section_banner(slide, title)

        # --- Image area (left side) ---
        # Priority: real image > auto-generated > gray placeholder > no image
        if image_prompt and not image_path:
            image_path = self._generate_image_for_slide(image_prompt, "content")
        has_image = False

        if image_path and os.path.exists(image_path):
            self._add_image(
                slide, image_path,
                left=Cm(2.5), top=Cm(5.5),
                max_width=Cm(9), max_height=Cm(9),
                name="img_content",
            )
            has_image = True
        elif image_placeholder:
            img_shape = self._add_shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                left=Cm(2.5),
                top=Cm(5.5),
                width=Cm(9),
                height=Cm(9),
                fill_color=RGBColor(0xE0, 0xE0, 0xE0),
                border_color=RGBColor(0xBD, 0xBD, 0xBD),
            )
            tf = img_shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = image_placeholder
            self._set_run_font(run, FONT_REGULAR, Pt(12), False, BODY_TEXT)
            has_image = True

        # --- Content body (with layout variants for visual variety) ---
        variant = self._content_layout_cycle % 3
        self._content_layout_cycle += 1

        content_top = Cm(5)
        content_height = Cm(11.5)

        if has_image:
            variant = 0
            content_left = Cm(13)
            content_width = Cm(18)
        else:
            content_left = Cm(3)
            content_width = Cm(28)

        if variant == 1 and not has_image:
            self._add_accent_stripe(slide)
            content_width = Cm(26)

        if variant == 2 and bullets and not has_image:
            self._add_numbered_points(slide, bullets, start_top=content_top + Cm(0.5))
        else:
            if bullets:
                card_shape = self._add_shape(
                    slide,
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    left=content_left - Cm(0.5),
                    top=content_top - Cm(0.3),
                    width=content_width + Cm(1),
                    height=content_height + Cm(0.6),
                    fill_color=CONTENT_CARD_BG,
                    border_color=CONTENT_CARD_BORDER,
                    border_width=Pt(1),
                    name="bg_content_card",
                    corner_radius=0.04,
                )
                self._add_shadow_to_shape(card_shape, blur_pt=4, opacity_pct=15)
                # Subtle gradient adds depth to content cards
                self._apply_gradient_fill(card_shape, "#F5F7FA", "#F0F2F5", angle=180)

                self._add_bullet_list(
                    slide,
                    left=content_left,
                    top=content_top,
                    width=content_width,
                    height=content_height,
                    items=bullets,
                    font_size=Pt(20),
                    name="txt_body",
                )
            elif paragraphs:
                card_shape = self._add_shape(
                    slide,
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    left=content_left - Cm(0.5),
                    top=content_top - Cm(0.3),
                    width=content_width + Cm(1),
                    height=content_height + Cm(0.6),
                    fill_color=CONTENT_CARD_BG,
                    border_color=CONTENT_CARD_BORDER,
                    border_width=Pt(1),
                    name="bg_content_card",
                    corner_radius=0.04,
                )
                self._add_shadow_to_shape(card_shape, blur_pt=4, opacity_pct=15)
                # Subtle gradient adds depth to content cards
                self._apply_gradient_fill(card_shape, "#F5F7FA", "#F0F2F5", angle=180)

                text = "\n\n".join(paragraphs)
                self._add_arabic_textbox(
                    slide,
                    left=content_left,
                    top=content_top,
                    width=content_width,
                    height=content_height,
                    text=text,
                    font_name=FONT_REGULAR,
                    font_size=Pt(18),
                    bold=False,
                    color=BODY_TEXT,
                    alignment=PP_ALIGN.RIGHT,
                    word_wrap=True,
                    auto_size=MSO_AUTO_SIZE.NONE,
                    name="txt_body",
                )

        # --- Speaker notes ---
        if notes:
            self._add_notes(slide, notes)

    def add_content_with_cards(
        self,
        title: str,
        cards: list,
        notes: str = "",
    ):
        """
        Add a content slide with 2-4 card layout.

        Cards are used for concepts where each card represents one
        concept with a title and optional body.

        Args:
            title: Section title
            cards: List of dicts with "title", "body" (optional),
                   "color" (optional), "image" (optional), "image_prompt" (optional)
            notes: Speaker notes / Storyline instructions
        """
        self.slide_count += 1
        slide = self._add_content_slide_with_layout()

        self._set_slide_title_for_toc(slide, title)
        self._add_header_bar(slide, self.lecture_title)
        self._add_section_banner(slide, title, wide=True)

        # --- Card layout ---
        card_count = len(cards)
        CARD_DARK2 = RGBColor(0x0E, 0x28, 0x41)
        AMBER = RGBColor(0xFF, 0x8F, 0x00)
        default_colors = [PRIMARY_BLUE, ACCENT1_BLUE, TEAL, CARD_DARK2, AMBER, PRIMARY_BLUE]

        cards_area_left = Cm(2.5)
        cards_area_width = Cm(28.5)
        cards_top = Cm(5.5)
        card_height = Cm(9)

        gap = Cm(0.8)
        total_gaps = gap * (card_count - 1) if card_count > 1 else 0
        card_width = int((cards_area_width - total_gaps) / card_count)

        for i, card_data in enumerate(cards):
            card_num = i + 1
            card_left = int(cards_area_left + i * (card_width + gap))
            card_color = card_data.get("color", default_colors[i % len(default_colors)])

            card_shape = self._add_shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                left=card_left,
                top=cards_top,
                width=card_width,
                height=card_height,
                fill_color=CARD_LIGHT_BG,
                border_color=card_color,
                border_width=Pt(2),
                name=f"card_{card_num}",
            )
            self._add_shadow_to_shape(card_shape)

            # Accent bar at top of card
            self._add_shape(
                slide,
                MSO_SHAPE.RECTANGLE,
                left=card_left,
                top=cards_top,
                width=card_width,
                height=Cm(1.2),
                fill_color=card_color,
            )

            # Optional card thumbnail image
            card_image = card_data.get("image")
            card_image_prompt = card_data.get("image_prompt")
            if card_image_prompt and not card_image:
                card_image = self._generate_image_for_slide(
                    card_image_prompt, "card",
                    topic_key=f"card_{i+1}" if not card_data.get("title") else None,
                )
            card_has_image = False
            if card_image:
                pic = self._add_image(
                    slide, card_image,
                    left=card_left + Cm(0.5),
                    top=cards_top + Cm(1.5),
                    max_width=card_width - Cm(1),
                    max_height=Cm(3),
                    name=f"img_card_{card_num}",
                )
                if pic is not None:
                    card_has_image = True

            title_top = cards_top + Cm(4.8) if card_has_image else cards_top + Cm(1.2)
            body_top = cards_top + Cm(6.5) if card_has_image else cards_top + Cm(3)
            body_height = Cm(2.5) if card_has_image else Cm(5.5)

            self._add_arabic_textbox(
                slide,
                left=card_left + Cm(0.5),
                top=title_top,
                width=card_width - Cm(1),
                height=Cm(1.5),
                text=card_data.get("title", ""),
                font_name=FONT_EXTRABOLD,
                font_size=Pt(20),
                bold=False,
                color=card_color if isinstance(card_color, RGBColor) else BODY_TEXT,
                alignment=PP_ALIGN.CENTER,
                name=f"txt_card_{card_num}_title",
            )

            body = card_data.get("body", "")
            if body:
                self._add_arabic_textbox(
                    slide,
                    left=card_left + Cm(0.5),
                    top=body_top,
                    width=card_width - Cm(1),
                    height=body_height,
                    text=body,
                    font_name=FONT_REGULAR,
                    font_size=Pt(18),
                    bold=False,
                    color=BODY_TEXT,
                    alignment=PP_ALIGN.RIGHT,
                    word_wrap=True,
                    auto_size=MSO_AUTO_SIZE.NONE,
                    name=f"txt_card_{card_num}_body",
                )

        if notes:
            self._add_notes(slide, notes)

    def add_two_column_slide(
        self,
        title: str,
        left_title: str,
        left_points: list,
        right_title: str,
        right_points: list,
        notes: str = "",
        right_image: Optional[str] = None,
        left_image: Optional[str] = None,
        right_image_prompt: Optional[str] = None,
        left_image_prompt: Optional[str] = None,
    ):
        """
        Add a two-column comparison slide.

        Args:
            title: Section title
            left_title: Title for the left column
            left_points: Bullet points for the left column
            right_title: Title for the right column
            right_points: Bullet points for the right column
            notes: Speaker notes
            right_image: Optional image above right column
            left_image: Optional image above left column
            right_image_prompt: Optional prompt to generate right image
            left_image_prompt: Optional prompt to generate left image
        """
        self.slide_count += 1
        slide = self._add_content_slide_with_layout()

        self._set_slide_title_for_toc(slide, title)
        self._add_header_bar(slide, self.lecture_title)
        self._add_section_banner(slide, title, wide=True)

        # Auto-generate column images from prompts
        if right_image_prompt and not right_image:
            right_image = self._generate_image_for_slide(right_image_prompt, "two_column")
        if left_image_prompt and not left_image:
            left_image = self._generate_image_for_slide(left_image_prompt, "two_column")

        # --- Column layout ---
        col_top = Cm(5)
        col_gap = Cm(1)
        col_width = Cm(13.5)

        right_col_left = Cm(17)
        left_col_left = Cm(2.5)

        right_has_img = right_image and os.path.exists(right_image)
        left_has_img = left_image and os.path.exists(left_image)
        any_has_img = right_has_img or left_has_img

        col_height = Cm(8) if any_has_img else Cm(10)
        img_shift = Cm(3.5)

        # --- Right column card ---
        right_card_height = col_height + (img_shift if right_has_img else 0)
        right_card = self._add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left=right_col_left - Cm(0.3),
            top=col_top - Cm(0.3),
            width=col_width + Cm(0.6),
            height=right_card_height + Cm(0.6),
            fill_color=WHITE,
            border_color=CONTENT_CARD_BORDER,
            border_width=Pt(1),
            name="bg_col1_card",
            corner_radius=0.04,
        )
        self._add_shadow_to_shape(right_card, blur_pt=4, opacity_pct=12)

        right_content_offset = 0
        if right_has_img:
            self._add_image(
                slide, right_image,
                left=right_col_left, top=col_top + Cm(0.3),
                max_width=col_width, max_height=Cm(3),
                name="img_col1",
            )
            right_content_offset = img_shift

        self._add_arabic_textbox(
            slide,
            left=right_col_left,
            top=col_top + right_content_offset,
            width=col_width,
            height=Cm(1.5),
            text=right_title,
            font_name=FONT_EXTRABOLD,
            font_size=Pt(20),
            bold=False,
            color=PRIMARY_BLUE,
            alignment=PP_ALIGN.CENTER,
            name="txt_col1_title",
        )

        self._add_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            left=right_col_left + Cm(2),
            top=col_top + Cm(1.5) + right_content_offset,
            width=col_width - Cm(4),
            height=Cm(0.15),
            fill_color=PRIMARY_BLUE,
        )

        self._add_bullet_list(
            slide,
            left=right_col_left,
            top=col_top + Cm(2) + right_content_offset,
            width=col_width,
            height=col_height - Cm(2),
            items=right_points,
            font_size=Pt(18),
            name="txt_col1_body",
        )

        # --- Left column card ---
        left_card_height = col_height + (img_shift if left_has_img else 0)
        left_card = self._add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left=left_col_left - Cm(0.3),
            top=col_top - Cm(0.3),
            width=col_width + Cm(0.6),
            height=left_card_height + Cm(0.6),
            fill_color=WHITE,
            border_color=CONTENT_CARD_BORDER,
            border_width=Pt(1),
            name="bg_col2_card",
            corner_radius=0.04,
        )
        self._add_shadow_to_shape(left_card, blur_pt=4, opacity_pct=12)

        left_content_offset = 0
        if left_has_img:
            self._add_image(
                slide, left_image,
                left=left_col_left, top=col_top + Cm(0.3),
                max_width=col_width, max_height=Cm(3),
                name="img_col2",
            )
            left_content_offset = img_shift

        self._add_arabic_textbox(
            slide,
            left=left_col_left,
            top=col_top + left_content_offset,
            width=col_width,
            height=Cm(1.5),
            text=left_title,
            font_name=FONT_EXTRABOLD,
            font_size=Pt(20),
            bold=False,
            color=ACCENT1_BLUE,
            alignment=PP_ALIGN.CENTER,
            name="txt_col2_title",
        )

        self._add_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            left=left_col_left + Cm(2),
            top=col_top + Cm(1.5) + left_content_offset,
            width=col_width - Cm(4),
            height=Cm(0.15),
            fill_color=ACCENT1_BLUE,
        )

        self._add_bullet_list(
            slide,
            left=left_col_left,
            top=col_top + Cm(2) + left_content_offset,
            width=col_width,
            height=col_height - Cm(2),
            items=left_points,
            font_size=Pt(18),
            name="txt_col2_body",
        )

        if notes:
            self._add_notes(slide, notes)
